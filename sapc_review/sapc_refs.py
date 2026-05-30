# -*- coding: utf-8 -*-
"""参考文献（GB/T 7714 风格；按撰写语言分外文/中文两组）"""
from sapc_common import add_heading_cn, add_body, set_cn_font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


# 外文文献（国外）
REFS_EN = [
    "Pecht M, et al. Failure of polymer aluminum electrolytic capacitors under elevated temperature humidity environments. CALCE, University of Maryland, 2017.（作者全名与期刊/卷期待核实）",
    "Sood B, Pecht M, et al. The effect of package geometry on moisture-driven degradation of polymer aluminum capacitors. M.S. Thesis / Report, University of Maryland (DRUM), 2016.（作者待核实）",
    "Teverovsky A. Degradation of ESR in polymer tantalum capacitors during high temperature storage. NASA NEPP Report, 2019. NTRS 20190030759.",
    "Teverovsky A. Breakdown and self-healing in tantalum capacitors. NASA NEPP Report, 2021. NTRS 20205008339.",
    "Teverovsky A. Reliability and breakdown of solid tantalum capacitors: field crystallization. NASA NEPP Reports, 2008-2021.（具体年代/编号待核实）",
    "Freeman Y, Lessner P, et al. Reliability of manganese dioxide and conductive polymer tantalum capacitors under temperature-humidity-bias testing. CARTS / Journal of the IMAPS (Allenpress), 2015.（作者待核实）",
    "KEMET. New reliability assessment practices for tantalum polymer capacitors. KEMET technical resources, 2020.",
    "KEMET. Advances in reliability of conducting polymer based capacitors in high humidity environment. KEMET technical resources, 2020.（指出分层为高湿下聚合物电容电导退化主因）",
    "Effect of high temperature storage on AC characteristics of polymer tantalum capacitors. IMAPS HiTEC Proceedings, 2021.（热氧化致电阻率随老化时间指数增长；作者待核实）",
    "Development of high temperature tantalum polymer capacitors. IMAPS / ISM Proceedings, 2018.（氧渗透路径与阻氧涂层；作者待核实）",
    "Thermal degradation mechanisms of PEDOT:PSS. Research report / journal article, 2008.（颗粒金属结构、导电畴收缩致老化；作者与出处待核实）",
    "Thermal stability investigation of PEDOT films from chemical oxidation and prepolymerized dispersion. 2013.（原位膜热稳定性弱于预聚合分散体膜；作者与出处待核实）",
    "Insight into the degradation mechanisms of highly conductive PEDOT thin films. 2020.（作者与期刊/卷期待核实）",
    "KyoceraAVX. Solid tantalum capacitors: MnO2 vs. polymer cathodes for high-reliability military and space applications. KyoceraAVX technical paper, 2021/2024.",
    "Murata Manufacturing. Conductive polymer aluminum electrolytic capacitors (ECAS): failure modes and mechanisms. Murata technical document.（开路为主市场失效、高温聚合物降解致容量偏移；版本/年代待核实）",
    "Panasonic. SP-Cap polymer aluminum capacitors (EEF series): layered aluminum technology, datasheets and reliability notes. Panasonic Industrial Devices.（系列与耐久参数随型号而异，待核实）",
    "Prokopowicz T I, Vaskas A R. Research and development, intrinsic reliability, subminiature ceramic capacitors. Final Report ECOM-90705-F, NTIS AD-864068, 1969.",
    "Waser R, Baiatu T, Hardtl K H. dc electrical degradation of perovskite-type titanates. Journal of the American Ceramic Society, 1990, 73(6): 1645-1673.",
    "Pozdeev-Freeman Y, Gill J. Crystallization of anodic Ta2O5 in solid tantalum capacitors. CARTS Proceedings, 2004.（卷期/页码待核实）",
    "Greason W D. Self-healing breakdown studies in metallized film capacitors. IEEE Transactions on Electrical Insulation / EOS, 1995.（卷期待核实）",
    "IEC 60384-1. Fixed capacitors for use in electronic equipment - Part 1: Generic specification. International Electrotechnical Commission.",
    "IEC 60384-24 / -25 / -26. Sectional specifications for surface mount fixed tantalum / aluminium electrolytic capacitors with conductive polymer solid electrolyte. IEC.（具体分部编号与年代待核实）",
    "AEC-Q200. Stress test qualification for passive components. Automotive Electronics Council.",
    "JEDEC JESD22 series. Reliability test methods for solid-state and passive devices. JEDEC Solid State Technology Association.",
    "US Patent Application 20240047142 A1. Method for preparing highly-reliable multilayer solid aluminum electrolytic capacitor. 2024.（叠层固态铝电解电容器结构与制法）",
]

# 中文文献（国内）
REFS_CN = [
    "弱碱去除 PEDOT:PSS 质子型掺杂以改善铝固态电解电容器性能的研究. RSC Advances, 2025.（中文团队成果，撰写语言/作者/卷期/DOI 待核实；若为英文撰写则应归入外文组）",
    "气相聚合 PEDOT 薄膜及其在固态电容器中的应用研究. 材料类期刊, 2021.（成膜结晶度对电导的影响；作者/卷期待核实）",
    "导电聚合物固态铝电解电容器材料与可靠性研究. 中文期刊/学位论文.（具体题名、作者、年份待核实）",
    "超薄层钛酸钡基贱金属电极 MLCC 的可靠性机理研究. 物理化学学报, 2024, 40(1): 2304015.（题名/卷期以原文为准，待核实）",
    "晶粒尺寸对多层陶瓷电容器可靠性的影响及其机理. 华南师范大学学报(自然科学版), 2024.（卷期/DOI 待核实）",
    "超薄 Ba(Ca)TiO3 基 MLCC 的介电温度稳定性与可靠性（中科院深圳先进电子材料研究院相关团队）. 中文/双语成果.（题名、作者、出处待核实）",
    "金属化膜电容器首次自愈的宏-微观分析. 高电压技术 / High Voltage, 2025.（撰写语言/作者/卷期待核实）",
    "金属化膜电容器自愈特性原位测试装置与分析方法. 科学仪器类期刊, 2024.（作者/卷期待核实）",
    "考虑自愈累积损伤的金属化聚丙烯薄膜电容器容量评估. 电子/电力电子类期刊, 2024.（作者/卷期待核实）",
    "基于超声方法的金属化膜电容器自愈放电检测. 电子类期刊, 2020.（作者/卷期待核实）",
    "固态/聚合物铝电解电容器湿热与高温退化及 ESR/容量演变研究. 中文期刊/学位论文.（题名、作者、年份待核实）",
    "车规级固态铝/聚合物电容器可靠性评估方法研究. 中文期刊/学位论文.（题名、作者、年份待核实）",
    "湖南艾华集团股份有限公司(上交所 603989). 固态聚合物/混合铝电解电容器产品与可靠性技术资料（低 ESR、良性失效\"不起火\"）. 企业技术资料.",
    "南通江海电容器股份有限公司(深交所 002484). 聚合物混合铝电解、薄膜及超级电容器产品与可靠性技术资料. 企业技术资料.",
    "一种高可靠性(多层)固态铝电解电容器的制造方法. 中国发明专利.（申请号/公开号、申请人、年份待核实；可与 US20240047142A1 对应核验）",
    "GB/T 2693《电子设备用固定电容器 第1部分：总规范》(等同采用 IEC 60384-1). 中国国家标准.（版本/年代待核实）",
    "GB/T 6346 系列《电子设备用固定电容器》分规范（含固态电解/聚合物电容相关分部）. 中国国家标准.（具体分部/年代待核实）",
    "GJB 360《电子及电气元件试验方法》及 GJB 系列高可靠电容器详细规范. 中国国家军用标准.（具体编号/年代待核实）",
]

NOTE = ("说明：本参考文献表按撰写语言分为\"外文文献\"（国外）与\"中文文献\"（国内）两组，"
        "以服务国内外研究现状的对照。凡标注\"待核实\"者，为本综述据公开资料整理但尚需在提交前"
        "于 CNKI / Web of Science / IEEE Xplore / 出版商页面 / 专利数据库逐条复核的卷期、DOI、"
        "作者或撰写语言信息，复核后请按 GB/T 7714—2015 统一著录；本综述不虚构任何文献、作者、年份、"
        "DOI 或数值，不确定处一律明示。若某条中文团队成果实际以英文发表，应据撰写语言改列入外文组。")


def _add_ref_list(doc, refs, start_idx):
    for i, ref in enumerate(refs, start_idx):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("[%d] %s" % (i, ref))
        set_cn_font(run, font_name="宋体", size_pt=10.5)
    return start_idx + len(refs)


def build_references(doc):
    add_heading_cn(doc, "参考文献", level=1)
    add_body(doc, NOTE, size=10.5)
    add_heading_cn(doc, "一、外文文献（国外）", level=3)
    nxt = _add_ref_list(doc, REFS_EN, 1)
    add_heading_cn(doc, "二、中文文献（国内）", level=3)
    _add_ref_list(doc, REFS_CN, nxt)
