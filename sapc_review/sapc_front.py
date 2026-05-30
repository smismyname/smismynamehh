# -*- coding: utf-8 -*-
"""前置：中文摘要、英文摘要、缩略语与术语表"""
from sapc_common import add_heading_cn, add_body, add_page_break, set_cn_font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def build_abstract(doc):
    add_heading_cn(doc, "摘  要", level=0)
    cn = """
叠层铝固态电容器（Stacked Aluminum Polymer Capacitor，SAPC；亦称导电聚合物固态铝电解电容器）以经蚀刻—阳极氧化的多片铝阳极箔并联叠层为正极、以导电聚合物聚（3,4-乙撑二氧噻吩）（PEDOT）替代传统液态电解质或二氧化锰作为固态阴极，兼具超低等效串联电阻（ESR）、优良的高频与纹波特性、以及"良性失效"（不燃烧、不爆浆）等突出优点，已成为高端电源、服务器、车载与通信电子中铝电解电容器与多层陶瓷电容器（MLCC）的重要替代与补充。然而，PEDOT 阴极对湿、热、电与机械应力的敏感性，使 SAPC 的失效行为与失效物理显著区别于其同族液态前身，并对其在严苛工况下的长期可靠性提出新的挑战。系统梳理 SAPC 失效模式与失效机理的国内外研究现状，是本学位论文绪论的核心任务。
本综述以 SAPC 为主线，以液态铝电解电容器、贱金属电极（BME）MLCC、金属化聚丙烯薄膜电容器、固体钽电容器（二氧化锰阴极与导电聚合物阴极）四类主流电容器为横向对照，沿"器件结构与材料体系→失效模式→失效机理"三模块的内在因果链组织，刻意避免"按时间罗列"与"国内国外两张皮"的低效写法，采用对照式、量化化的系统评述，并将国内（中文文献）与国外（外文文献）进展在同一主题下并置比较。综述采用"主题数据库检索＋引文滚雪球＋权威机构工程报告"三层文献获取策略，时间跨度自20世纪经典电解/陶瓷电容失效物理研究至2025年最新的原位表征与多物理场机理研究，并对 PEDOT 阴极体系及2021—2025近五年文献重点采纳。
在器件模块，本综述厘清 SAPC 的叠层结构、PEDOT 阴极体系（原位化学氧化聚合、气相聚合、PEDOT:PSS 浸渍三条工艺路线）及其工艺演化，并与四类对照器件在介质体系、阴极/电解质、ESR与纹波能力、电压与温度范围等维度展开结构对比；在失效模式模块，归纳 SAPC 以 ESR 上升与容量下降（磨损型、以开路为主导市场失效）、漏电流上升乃至介质击穿短路（偶发、良性）为主的失效模式谱，并与四类器件的开路—短路—参数漂移三分法对照；在失效机理模块，深入评述 PEDOT 去掺杂与热氧化（导电晶粒收缩、跳跃输运、氧渗透）、氧化铝介质水合、界面分层（热-机-电耦合）以及温-电-湿多应力耦合等核心机理，并以激活能 Ea、湿度指数、退化速率等量化参数与液态铝电解的电解液蒸发、MLCC 的氧空位迁移、薄膜的自愈累积、钽的场致结晶与热失控等机理进行横向比较。
通过三模块的量化对照与文献计量分析，本综述得出总体判断：国外（尤以 NASA NEPP、KEMET、Panasonic、Nichicon、Nippon Chemi-Con 等机构与厂商）在 PEDOT 机理认识、原位表征与工程判据方面长期引领，国内（中文文献）近年在国产材料工艺、车规可靠性与数据驱动评估方面快速跟进、工程落地见长，但在原创机理理论、强物理嵌入与跨工况泛化方面仍存差距。在此基础上，本综述凝练出四条研究空白——SAPC 失效率（FIT）基准缺位、ESR/容量/漏电失效阈值判据不统一、温-电-湿多应力耦合机理缺乏定量理论、PEDOT 去掺杂/热氧化动力学未被定量刻画，并逐条对应到学位论文的研究内容与章节安排，最后展望了稳定化 PEDOT 材料、原位多物理场机理、强物理嵌入混合建模与自主化可靠性保证等未来趋势。
"""
    add_body(doc, cn)
    add_heading_cn(doc, "关键词：", level=4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0.82)
    run = p.add_run("叠层铝固态电容器；导电聚合物；PEDOT；失效模式；失效机理；去掺杂；"
                    "热氧化；电化学迁移；分层；ESR 退化；湿热偏置；国内外研究现状")
    set_cn_font(run, font_name="宋体", size_pt=11)
    add_page_break(doc)

    add_heading_cn(doc, "Abstract", level=0)
    en = """
The stacked aluminum polymer capacitor (SAPC, also called the conductive-polymer solid aluminum electrolytic capacitor) employs multiple parallel-laminated etched-and-anodized aluminum anode foils as the anode and a conductive polymer, poly(3,4-ethylenedioxythiophene) (PEDOT), as the solid cathode in place of the conventional liquid electrolyte or manganese dioxide. It combines ultra-low equivalent series resistance (ESR), excellent high-frequency and ripple-current behavior, and a benign (non-ignition) failure response, and has become an important replacement for and complement to aluminum electrolytic capacitors and multilayer ceramic capacitors (MLCC) in high-end power supplies, servers, automotive and communication electronics. However, the sensitivity of the PEDOT cathode to humidity, heat, electrical and mechanical stresses makes the failure behavior and failure physics of SAPC markedly different from those of its liquid-electrolyte predecessor, posing new challenges for long-term reliability under harsh missions. A systematic review of the domestic and international research status of SAPC failure modes and mechanisms is therefore the core task of the introduction chapter of this dissertation.
Taking SAPC as the main line and four mainstream families - liquid aluminum electrolytic, base-metal-electrode (BME) MLCC, metallized polypropylene film, and solid tantalum (MnO2 and conductive-polymer cathodes) - as lateral comparisons, this review is organized along the intrinsic causal chain of three modules: device structure & materials, failure modes, and failure mechanisms. It deliberately avoids chronological listing and the separation of domestic and international work into two disjoint parts, adopting a contrastive and quantitative style in which Chinese-language (domestic) and foreign-language (international) progress are juxtaposed under the same topic. A three-layer literature-acquisition strategy (database retrieval, citation snowballing, and authoritative engineering reports) is used, spanning from twentieth-century classical electrolytic/ceramic failure physics to the latest in-situ characterization and multi-physics mechanism studies in 2025, with emphasis on the PEDOT cathode system and 2021-2025 literature.
At the device level, the review clarifies the stacked structure of SAPC, its PEDOT cathode systems (in-situ chemical oxidative polymerization, vapor-phase polymerization, and PEDOT:PSS dispersion impregnation) and process evolution, and compares them with the four reference families. At the failure-mode level, it summarizes the SAPC mode spectrum dominated by ESR rise and capacitance drop (wear-out, open mode dominating field failures) and leakage rise to dielectric-breakdown short (accidental, benign), and contrasts it with the open-short-drift taxonomy of the four families. At the mechanism level, it reviews in depth PEDOT de-doping and thermo-oxidation (conductive-grain shrinkage, hopping transport, oxygen permeation), hydration of the alumina dielectric, interface delamination (thermo-mechanical-electrical coupling), and temperature-voltage-humidity multi-stress coupling, comparing them quantitatively (activation energy, humidity exponent, degradation rate) with electrolyte evaporation in liquid aluminum, oxygen-vacancy migration in MLCC, self-healing accumulation in film, and field crystallization and thermal runaway in tantalum.
Through quantitative cross-comparison and bibliometric analysis of the three modules, the review concludes that international groups (NASA NEPP, KEMET, Panasonic, Nichicon, Nippon Chemi-Con, etc.) lead in PEDOT mechanism understanding, in-situ characterization and engineering criteria, while domestic (Chinese-language) work has advanced rapidly in localized materials/processes, automotive-grade reliability and data-driven assessment, yet gaps remain in original mechanistic theory, strongly physics-embedded modeling and cross-condition generalization. On this basis, four research gaps are distilled - the absence of an SAPC failure-rate (FIT) benchmark, inconsistent end-of-life criteria for ESR/capacitance/leakage, the lack of a quantitative theory for temperature-voltage-humidity multi-stress coupling, and the un-quantified kinetics of PEDOT de-doping/thermo-oxidation - and mapped onto the dissertation research contents, followed by an outlook on stabilized PEDOT materials, in-situ multi-physics mechanism study, strongly physics-embedded hybrid modeling and autonomous reliability assurance.

Keywords: stacked aluminum polymer capacitor; conductive polymer; PEDOT; failure mode; failure mechanism; de-doping; thermo-oxidation; electrochemical migration; delamination; ESR degradation; temperature-humidity-bias; research status.
"""
    add_body(doc, en)
    add_page_break(doc)


def build_glossary(doc):
    """缩略语与主要术语表（不使用编号表格，避免影响章内表号）"""
    add_heading_cn(doc, "缩略语与主要术语", level=0)
    add_body(doc, """
为便于阅读，现将本综述频繁出现的缩略语与主要术语集中说明如下（按英文字母与主题归类）。正文首次出现时一般给出中英文对照，此处提供统一索引。其中标注"待核实"者，为本综述据公开资料整理但尚需在提交前按原始文献逐条复核的数值或表述。
""")
    items = [
        ("SAPC", "Stacked Aluminum Polymer Capacitor，叠层铝固态（导电聚合物）电容器，本综述主线对象"),
        ("PAEC / PA-cap", "Polymer Aluminum Electrolytic Capacitor，聚合物铝电解电容器（SAPC 的等义/近义表述）"),
        ("PEDOT", "Poly(3,4-ethylenedioxythiophene)，聚（3,4-乙撑二氧噻吩），SAPC 的导电聚合物阴极主体"),
        ("PEDOT:PSS", "PEDOT 与聚苯乙烯磺酸（PSS）的复合分散体，常用阴极成型材料；PSS 兼作掺杂剂/反离子"),
        ("PSS", "Poly(styrene sulfonate)，聚苯乙烯磺酸（根），强酸性、吸湿性高分子掺杂剂"),
        ("VPP", "Vapor-Phase Polymerization，气相聚合（PEDOT 成膜工艺之一）"),
        ("MnO2", "二氧化锰，传统固态钽/铝电容阴极材料，是浪涌点燃的氧来源"),
        ("MLCC", "Multilayer Ceramic Capacitor，多层陶瓷电容器（对照器件）"),
        ("BME", "Base-Metal-Electrode，贱金属电极（镍内电极、铜端电极 MLCC 工艺）"),
        ("BOPP", "Biaxially-Oriented PolyPropylene，双向拉伸聚丙烯（薄膜电容介质）"),
        ("ESR", "Equivalent Series Resistance，等效串联电阻，SAPC 的核心健康/退化指标"),
        ("ESL", "Equivalent Series Inductance，等效串联电感"),
        ("C", "Capacitance，电容量"),
        ("DCL / LC", "DC Leakage Current，直流漏电流（聚合物/钽电容核心退化量之一）"),
        ("tanδ", "Loss tangent，损耗角正切（介质/整体损耗指标）"),
        ("IR", "Insulation Resistance，绝缘电阻（MLCC 核心退化量）"),
        ("THB / TH", "Temperature-Humidity-Bias / Temperature-Humidity，温湿度偏压/温湿度试验"),
        ("HAST", "Highly Accelerated temperature & humidity Stress Test，高加速温湿度应力试验"),
        ("HTS", "High-Temperature Storage，高温贮存（聚合物热氧化研究常用）"),
        ("TC", "Temperature Cycling，温度循环"),
        ("Reflow", "回流焊，表贴器件经历的高温（约260℃无铅）焊接热冲击"),
        ("CTE", "Coefficient of Thermal Expansion，热膨胀系数（分层/裂纹的关键参量）"),
        ("EoL", "End of Life，寿命终点（失效阈值判据，如 ESR 翻倍、容量下降10%等）"),
        ("De-doping", "去掺杂，导电聚合物失去掺杂态正电荷载流子、电导率下降的过程"),
        ("Thermo-oxidation", "热氧化，氧在热作用下使 PEDOT 共轭结构降解、导电晶粒收缩的过程"),
        ("Hydration", "水合/水解，阳极氧化铝（Al2O3）与水生成 AlOOH/Al(OH)3 的过程"),
        ("Delamination", "分层，阴极聚合物—碳—银—引线框等界面在热-机-电应力下脱开"),
        ("Self-healing", "自愈，金属化薄膜电容缺陷点电极蒸发隔离机制；聚合物阴极亦有类自愈隔离"),
        ("Field crystallization", "场致结晶，钽电容 Ta2O5 非晶向晶态转变，致漏电激增"),
        ("ECM", "Electrochemical Migration，电化学迁移（湿热偏置下金属离子迁移成枝）"),
        ("Ea", "Activation energy，激活能（温度敏感性参数，单位 eV）"),
        ("FIT", "Failures In Time，每 10^9 器件·小时的失效数（失效率单位）"),
        ("β / η", "Weibull 分布的形状参数 / 尺度（特征寿命）参数"),
        ("PoF", "Physics-of-Failure，失效物理"),
        ("FMEA / FMMEA", "失效模式（机理）与影响分析"),
        ("NEPP", "NASA Electronic Parts and Packaging Program，美国航空航天局电子元器件与封装可靠性项目"),
        ("待核实", "本综述据公开资料整理、但需按原始文献核验的具体数值/卷期/DOI/作者信息"),
    ]
    for k, v in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(2.4)
        p.paragraph_format.first_line_indent = Cm(-2.4)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("%s　" % k)
        set_cn_font(run, font_name="黑体", size_pt=10.5, bold=True)
        run2 = p.add_run(v)
        set_cn_font(run2, font_name="宋体", size_pt=10.5)
    add_page_break(doc)
