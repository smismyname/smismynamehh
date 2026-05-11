# -*- coding: utf-8 -*-
"""Word 文档构建工具函数 (SAPC Review)"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FIG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")


def set_cn_font(run, font_name="宋体", size_pt=11, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), "Times New Roman")
    rFonts.set(qn('w:hAnsi'), "Times New Roman")


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc, text, level=1):
    sizes = {0: 22, 1: 18, 2: 14, 3: 12.5, 4: 11.5}
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_cn_font(run, font_name="黑体", size_pt=sizes.get(level, 11.5), bold=True)
    pPr = p._p.get_or_add_pPr()
    oLvl = OxmlElement('w:outlineLvl')
    oLvl.set(qn('w:val'), str(level))
    pPr.append(oLvl)
    return p


def add_body(doc, text, indent=True, size=10.5):
    """支持多段正文输入 (用\n\n分段)"""
    paragraphs = [s.strip() for s in text.strip().split('\n\n') if s.strip()]
    for para_text in paragraphs:
        para_text = ' '.join([ln.strip() for ln in para_text.split('\n') if ln.strip()])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(para_text)
        set_cn_font(run, font_name="宋体", size_pt=size)


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_cn_font(run, font_name="Times New Roman", size_pt=11)


def add_figure(doc, filename, caption, width_cm=14.5):
    """插入图片+图注"""
    path = os.path.join(FIG_DIR, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(6)
    run2 = p2.add_run(caption)
    set_cn_font(run2, font_name="黑体", size_pt=10, bold=True)


def add_table(doc, headers, rows, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        set_cn_font(run, font_name="黑体", size_pt=10, bold=True)
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tbl.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(h)
        set_cn_font(r, font_name="黑体", size_pt=9.5, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri+1].cells[ci]
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = cp.add_run(str(val))
            set_cn_font(r, font_name="宋体", size_pt=9.5)
    return tbl


def add_quote(doc, text, size=10):
    """缩进引述段落"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_cn_font(run, font_name="楷体", size_pt=size)


def add_reference_item(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    set_cn_font(run, font_name="宋体", size_pt=9.5)


def configure_page(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for xml in [
        ('fldChar', {'fldCharType': 'begin'}),
    ]:
        pass
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = "PAGE"
    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar_begin)
    run._r.append(instrText)
    run._r.append(fldChar_end)
    set_cn_font(run, font_name="Times New Roman", size_pt=10)
