# -*- coding: utf-8 -*-
"""参考文献（按层级与主题归类；部分条目卷期/页码/DOI标注“建议核验”）"""
from sapc_common import add_heading_cn, set_cn_font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


REFS = [
    # ===== 综述/方法学与电容可靠性总论 =====
    "Wang H, Blaabjerg F. Reliability of capacitors for DC-link applications in power electronic converters—An overview. IEEE Transactions on Industry Applications, 2014, 50(5): 3569-3578.",
    "Wang H, Liserre M, Blaabjerg F. Toward reliable power electronics: Challenges, design tools, and opportunities. IEEE Industrial Electronics Magazine, 2013, 7(2): 17-26.",
    "Zhao Z, Davari P, Wang H, et al. An overview of condition monitoring techniques for capacitors in DC-link applications. IEEE Transactions on Power Electronics, 2021, 36(4): 3692-3716.（卷期建议核验）",
    "Yang S, Bryant A, Mawby P, et al. An industry-based survey of reliability in power electronic converters. IEEE Transactions on Industry Applications, 2011, 47(3): 1441-1451.",
    "Choi U M, Blaabjerg F, et al. Power electronics reliability: state of the art and outlook. IEEE Journal of Emerging and Selected Topics in Power Electronics, 2023.（卷期/页码建议核验）",

    # ===== 导电高分子与PEDOT材料 =====
    "Groenendaal L, Jonas F, Freitag D, et al. Poly(3,4-ethylenedioxythiophene) and its derivatives: past, present, and future. Advanced Materials, 2000, 12(7): 481-494.",
    "Elschner A, Kirchmeyer S, Lovenich W, et al. PEDOT: Principles and Applications of an Intrinsically Conductive Polymer. CRC Press, 2010.",
    "Shi H, Liu C, Jiang Q, et al. Effective approaches to improve the electrical conductivity of PEDOT:PSS: a review. Advanced Electronic Materials, 2015, 1(4): 1500017.",
    "Stretchable conductive polymers and composites based on PEDOT and PEDOT:PSS (review). 2019, PMC6401235.（作者/卷期建议核验）",

    # ===== SAPC / 导电高分子电解电容结构与机理 =====
    "Panasonic Industrial Devices. SP-Cap Conductive Polymer Aluminium Electrolytic Capacitors — product and technology documentation (EEF-GX/SX/JX/KX series). 技术资料.",
    "Panasonic Industrial Devices. Conductive Polymer Aluminium Solid Capacitors (OS-CON) — technology documentation. 技术资料.",
    "Nichicon Corporation. Latest technological trends in conductive polymer aluminium solid electrolytic capacitors (technical article). 2019.",
    "Nichicon Corporation. Conductive polymer hybrid aluminium electrolytic capacitor (GYA/GYB/GYC/GYD) white paper. 2020.",
    "Murata Manufacturing. Polymer aluminium electrolytic capacitor (ECAS series) failure-mode and reliability documentation. 技术资料.",
    "Freeman Y, Lessner P, et al. Reliability and reliability physics of polymer tantalum and aluminium capacitors. KEMET technical papers / CARTS Proceedings.（年代/卷期建议核验）",
    "KEMET. Advances in reliability of conducting polymer based capacitors in high humidity environment (technical paper). 2019.",

    # ===== 失效分析（CALCE/NASA等） =====
    "Nishikawa/CALCE (University of Maryland). Reliability evaluation of liquid and polymer aluminium electrolytic capacitors (M.S. thesis). 2014.（作者姓名建议核验）",
    "CALCE EPSC. Failure of polymer aluminium electrolytic capacitors under elevated temperature-humidity environments. 2014-2017.（出处/年份建议核验）",
    "Teverovsky A. Breakdown and self-healing in tantalum capacitors. NASA NEPP Report, 2021 (NTRS 20205008339).",
    "Teverovsky A. Effect of moisture on leakage currents and breakdown in solid polymer/tantalum capacitors. NASA NEPP Report.（年份建议核验）",
    "Reversible post-breakdown conduction in aluminium oxide-polymer capacitors. Applied Physics Letters, 2013, 102(16): 163504, DOI: 10.1063/1.4802485.（作者建议核验）",

    # ===== 工艺/分散液/专利 =====
    "Solid electrolytic capacitor containing an intrinsically conductive polymer. US Patent Application 20230026663, 2023.",
    "Solid electrolytic capacitor containing a moisture barrier. US Patent Application 20210383980, 2021.",
    "Method of manufacturing solid electrolytic capacitor and dispersion containing conjugated electrically conductive polymer. US Patent Application 20210027953, 2021.",
    "PSS dispersions and aluminium electrodes for solid-state electrolytic capacitors. Inorganics (MDPI), 2024, 12(4): 104.（作者建议核验）",

    # ===== 可靠性建模/寿命/加速 =====
    "Parler S G. Deriving life multipliers for electrolytic capacitors. IEEE Power Electronics Society Newsletter, 2004, 16(1): 11-12.",
    "Gasperi M L. Life prediction modeling of bus capacitors in AC variable-frequency drives. IEEE Transactions on Industry Applications, 2005, 41(6): 1430-1435.",
    "Lifetime prediction and reliability analysis for aluminium electrolytic capacitors in EV charging module based on mission profiles. Frontiers in Electronics, 2023, 4: 1226006, DOI: 10.3389/felec.2023.1226006.",
    "Kim M, et al. Physics-constrained / physics-informed machine learning for failure-time prediction of capacitors (MLCC). APL Machine Learning, 2023.（卷期建议核验）",

    # ===== 状态监测/RUL（含国内） =====
    "Lu W G, Zhou L W, Du X, et al. Grey-box online monitoring method for DC-link capacitors (Chongqing University). IEEE Transactions on Power Electronics / 电工技术学报, 卷期与年份建议核验.",
    "Jian Y, Chen Z, Peng S, et al. Capacitor aging state evaluation and a remaining-useful-life prediction method based on a CNN-LSTM network considering the impact of parameter dispersion. Electronics (MDPI), 2025, 14(22): 4452.",
    "Huang Z, et al. An online remaining-useful-life prediction method for tantalum capacitors based on temperature measurements. Electronics (MDPI), 2025, 14(22): 4393.",
    "A physics-informed neural network-based method for predicting degradation trajectories and remaining useful life of supercapacitors. Green Energy and Intelligent Transportation, 2025, DOI: 10.1016/j.geits.2025.100291.",
    "Composite denoising-based LSTM prediction method of supercapacitor performance degradation and remaining useful life. Circuit World (Emerald), 2024, DOI: 10.1108/CW-12-2023-0459.（作者建议核验）",
    "Song W, Yang X, Deng W, et al. Remaining useful life prediction for power storage electronic components based on fractional Weibull process and shock Poisson model. Fractal and Fractional (MDPI), 2024, 8(8): 485.",

    # ===== 国内材料/工艺（PEDOT:PSS改性、气相聚合、固态铝聚合物） =====
    "Removal of protonic doping from PEDOT:PSS by weak base for improving aluminium solid electrolytic capacitor performance. RSC Advances, 2025, DOI: 10.1039/D5RA00124B.（作者与卷期建议核验）",
    "High-performance solid capacitor using vapor-phase-polymerized PEDOT film. Journal of Materials Science, 2021, 56, DOI: 10.1007/s10853-021-05976-1.（作者建议核验）",
    "Liquid electrolyte-free cylindrical aluminium polymer capacitor: materials and characteristics (review). 2015.（卷期/作者建议核验）",
    "高压电解电容器阳极铝箔氧化膜微观结构演变. Journal of Alloys and Compounds, 2020.（卷期/作者建议核验）",

    # ===== 相邻领域（金属化膜自愈/钽，供机理借鉴） =====
    "Unveiling first self-healing in metallised film capacitor: a macro-micro analysis. High Voltage (CSEE/IET), 2025.（DOI/卷期建议核验）",
    "Study of the in-situ test setup and analysis methods for self-healing properties of metallised film capacitors. Review of Scientific Instruments, 2024, 95(4): 045105, DOI: 10.1063/5.0194057.",
    "Capacitance evaluation of metallized polypropylene film capacitors considering cumulative self-healing damage. Electronics (MDPI), 2024, 13(14): 2886.",
    "Ma Y, et al. Detection of self-healing discharge in metallized film capacitors using an ultrasonic method. Electronics (MDPI), 2020, 9(11): 1893.",

    # ===== 标准/规范 =====
    "IEC 60384-1. Fixed capacitors for use in electronic equipment - Part 1: Generic specification. International Electrotechnical Commission.",
    "IEC 60384-24/25/26. Sectional specifications for surface-mount and conductive-polymer solid electrolytic capacitors. IEC.（具体分部/年代建议核验）",
    "AEC-Q200. Stress test qualification for passive components. Automotive Electronics Council.",
    "JEDEC J-STD-020. Moisture/reflow sensitivity classification for surface-mount devices. JEDEC.",
    "GB/T 2693《电子设备用固定电容器 第1部分：总规范》(等同采用 IEC 60384-1). 中国国家标准（年代/版本建议核验）.",
    "GJB 360《电子及电气元件试验方法》及 GJB 系列高可靠电容器详细规范. 中国国家军用标准（具体编号/年代建议核验）.",

    # ===== 市场/产业 =====
    "Astute Analytica. Conductive polymer capacitor market (global / Japan) trends and competition analysis, 2025-2035. 2026.（市场口径仅供定性参考）",
    "Precedence Research. Conductive polymer capacitor market size and forecast. 2025.（市场口径仅供定性参考）",
    "艾华集团股份有限公司（湖南，上交所 603989）. 固态聚合物/多层聚合物铝电解电容器产品与可靠性技术资料（低ESR、不起火良性失效）. 企业技术资料.",
    "南通江海电容器股份有限公司（深交所 002484）. 铝电解、薄膜、聚合物及超级电容器产品与可靠性技术资料. 企业技术资料.",

    # ===== 说明 =====
    "说明：本报告SAPC研究现状部分在上述可溯源公开文献、厂商技术资料与市场报告的基础上撰写，并综合参考了国内外在导电高分子材料、固态铝/聚合物电容工艺、失效机理、寿命建模与在线监测等方向的公开成果；受检索条件限制，凡未能确证的作者、卷期、页码或DOI均已标注“建议核验”，提交前请在 CNKI / Web of Science / IEEE Xplore / Scopus 及出版商页面逐条复核，并按 GB/T 7714 统一著录。市场规模数据来自第三方研究机构，不同机构口径差异较大，正文仅取其增速量级作定性判据。",
]


def build_references(doc):
    add_heading_cn(doc, "参考文献", level=1)
    for i, ref in enumerate(REFS, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("[%d] %s" % (i, ref))
        set_cn_font(run, font_name="宋体", size_pt=10.5)
