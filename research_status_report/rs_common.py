# -*- coding: utf-8 -*-
"""
研究现状专题报告 - 通用样式与排版辅助模块
Common styling / layout helpers for the "Capacitor Reliability Research Status" report.

设计目标：
1. 自包含，不依赖 report_generator 包，便于独立构建与下载。
2. 提供中文标题/正文/公式/三线表/插图/自动目录(域)等排版能力。
3. 图内统一英文/数字标注（沙箱无 CJK 字体），图题(中文)由 Word 渲染。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ----------------------------------------------------------------------------
# 全局计数器（图、表、公式）
# ----------------------------------------------------------------------------
class Counters:
    def __init__(self):
        self.fig = 0
        self.tab = 0
        self.eq = 0

    def next_fig(self):
        self.fig += 1
        return self.fig

    def next_tab(self):
        self.tab += 1
        return self.tab

    def next_eq(self):
        self.eq += 1
        return self.eq


COUNT = Counters()


# ----------------------------------------------------------------------------
# 字体
# ----------------------------------------------------------------------------
def set_cn_font(run, font_name="宋体", size_pt=12, bold=False, color=None, italic=False):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), "Times New Roman")
    rFonts.set(qn('w:hAnsi'), "Times New Roman")


# ----------------------------------------------------------------------------
# 分页 / 页面
# ----------------------------------------------------------------------------
def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def configure_page(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = "PAGE"
    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar_begin)
    run._r.append(instrText)
    run._r.append(fldChar_end)
    set_cn_font(run, font_name="Times New Roman", size_pt=10)


def add_header_text(doc, text):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_cn_font(run, font_name="宋体", size_pt=9)
    # 页眉下边框
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '808080')
    pBdr.append(bottom)
    pPr.append(pBdr)


# ----------------------------------------------------------------------------
# 标题
# ----------------------------------------------------------------------------
def add_heading_cn(doc, text, level=1):
    sizes = {0: 22, 1: 16, 2: 14, 3: 12.5, 4: 12}
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_cn_font(run, font_name="黑体", size_pt=sizes.get(level, 12), bold=True)
    p.paragraph_format.space_before = Pt(10 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    outlineLvl = OxmlElement('w:outlineLvl')
    outlineLvl.set(qn('w:val'), str(level))
    pPr.append(outlineLvl)
    return p


# ----------------------------------------------------------------------------
# 正文
# ----------------------------------------------------------------------------
def add_body(doc, text, first_indent=True, size=12, after=4):
    paras = []
    for para_text in text.strip().split('\n'):
        para_text = para_text.strip()
        if not para_text:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first_indent:
            p.paragraph_format.first_line_indent = Cm(0.82)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(para_text)
        set_cn_font(run, font_name="宋体", size_pt=size)
        paras.append(p)
    return paras


def add_quote(doc, text):
    """评述 / 重点框：左缩进、楷体、带左边框"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_cn_font(run, font_name="楷体", size_pt=11.5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '6')
    left.set(qn('w:color'), '4472C4')
    pBdr.append(left)
    pPr.append(pBdr)
    return p


def add_formula(doc, text, number=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_cn_font(run, font_name="Times New Roman", size_pt=11.5, italic=False)
    if number:
        n = COUNT.next_eq()
        tab = p.add_run("        （%d）" % n)
        set_cn_font(tab, font_name="Times New Roman", size_pt=11)
    return p


# ----------------------------------------------------------------------------
# 三线表
# ----------------------------------------------------------------------------
def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'bottom', 'left', 'right'):
        if edge in kwargs:
            spec = kwargs[edge]
            el = tcBorders.find(qn('w:%s' % edge))
            if el is None:
                el = OxmlElement('w:%s' % edge)
                tcBorders.append(el)
            el.set(qn('w:val'), spec.get('val', 'single'))
            el.set(qn('w:sz'), str(spec.get('sz', 8)))
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), spec.get('color', '000000'))


def add_table(doc, headers, rows, caption=None, col_widths=None, font_size=9.5,
              header_size=9.5):
    """三线表：上下粗线、表头下细线。caption 自动编号。"""
    if caption:
        n = COUNT.next_tab()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run("表 %d  %s" % (n, caption))
        set_cn_font(run, font_name="黑体", size_pt=10.5, bold=True)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True

    thick = {'val': 'single', 'sz': 14, 'color': '000000'}
    thin = {'val': 'single', 'sz': 6, 'color': '000000'}
    nrows = 1 + len(rows)

    # 表头
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(str(h))
        set_cn_font(run, font_name="黑体", size_pt=header_size, bold=True)
        _set_cell_border(hdr[i], top=thick, bottom=thin)

    # 数据
    for r_idx, row in enumerate(rows):
        is_last = (r_idx == len(rows) - 1)
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(str(val))
            set_cn_font(run, font_name="宋体", size_pt=font_size)
            if is_last:
                _set_cell_border(cell, bottom=thick)

    if col_widths:
        for i, w in enumerate(col_widths):
            for r in range(nrows):
                table.rows[r].cells[i].width = Cm(w)
    return table


# ----------------------------------------------------------------------------
# 插图
# ----------------------------------------------------------------------------
def add_figure(doc, img_path, caption, width_cm=14.5):
    if not os.path.exists(img_path):
        # 占位
        add_body(doc, "[图缺失: %s]" % img_path)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    n = COUNT.next_fig()
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    crun = cap.add_run("图 %d  %s" % (n, caption))
    set_cn_font(crun, font_name="黑体", size_pt=10.5, bold=True)
    return n


# ----------------------------------------------------------------------------
# 列表项（带编号/项目符号的正文）
# ----------------------------------------------------------------------------
def add_list_item(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.82)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_cn_font(run, font_name="宋体", size_pt=size)
    return p


# ----------------------------------------------------------------------------
# 自动目录（TOC 域）
# ----------------------------------------------------------------------------
def add_toc_field(doc):
    add_heading_cn(doc, "目  录", level=0)
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar_sep = OxmlElement('w:fldChar')
    fldChar_sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = "（请在 Word 中按 Ctrl+A 后按 F9 更新目录域以生成页码）"
    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar)
    r.append(instrText)
    r.append(fldChar_sep)
    r.append(t)
    r.append(fldChar_end)
    set_cn_font(run, font_name="宋体", size_pt=11)
    add_page_break(doc)


# ----------------------------------------------------------------------------
# 封面
# ----------------------------------------------------------------------------
def build_cover(doc, title_lines, subtitle_en, meta):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("专  题  综  述  报  告")
    set_cn_font(run, font_name="黑体", size_pt=26, bold=True)
    doc.add_paragraph()
    doc.add_paragraph()
    for line in title_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_cn_font(run, font_name="黑体", size_pt=22, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle_en)
    set_cn_font(run, font_name="Times New Roman", size_pt=13, italic=True)
    for _ in range(6):
        doc.add_paragraph()
    for k, v in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("%s：%s" % (k, v))
        set_cn_font(run, font_name="宋体", size_pt=14)
    add_page_break(doc)


def new_document():
    doc = Document()
    configure_page(doc)
    add_page_number_footer(doc)
    add_header_text(doc, "电容器可靠性研究现状专题综述")
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc
