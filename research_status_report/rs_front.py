# -*- coding: utf-8 -*-
"""前置：摘要（中/英）"""
from rs_common import add_heading_cn, add_body, add_page_break, set_cn_font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def build_abstract(doc):
    add_heading_cn(doc, "摘  要", level=0)
    cn = """
电容器是电力电子系统与现代信息装备中用量最大、最为基础的无源储能元件，也常是整机系统中故障率最高、寿命最短的"薄弱环节"。系统梳理电容器可靠性的国内外研究现状，是电子元器件可靠性工程与博士学位论文绪论的核心任务。本报告以铝电解、多层陶瓷（MLCC）、金属化薄膜、固体钽与超级电容五大类电容器为对象，刻意避免"按时间罗列"与"国内国外两张皮"的低效写法，沿"器件—失效模式—失效机理—可靠性建模—状态监测"五环节主线，对国内外研究进行对照式、量化化的系统评述。
报告首先建立综述的范围、文献检索策略与组织主线（立坐标）；继而沿主线逐环节展开：在器件环节确立电容器作为系统可靠性短板的故障归因地位（直流母线电容约占变流器故障30%，电解电容21%–60%）；在失效模式环节梳理开路—短路—参数漂移三分法与五类器件失效模式谱；在失效机理环节深入评述电解液蒸发（Ea≈0.94 eV）、氧空位迁移（Ea≈0.88–1.49 eV）、自愈累积（终点10^4–10^6次、容量降5%–10%）、场致结晶（电压指数n≈15–40）等核心机理；在可靠性建模环节对比统计、加速/Prokopowicz-Vaskas、广义Eyring、失效物理与数据驱动/PINN混合五类方法；在状态监测环节对比ESR/C在线辨识（精度自C±4%/ESR±12%提升至C±1.5%/ESR±5%）与机器学习RUL预测（RMSE自约12%降至约3.8%）的国内外进展。
通过综合量化对比、文献计量、技术成熟度与研究演进时间轴的多视角分析，报告得出总体判断：国内外呈"国外引领范式、国内深化工程"的互补格局，差距集中在原创理论、强物理嵌入建模与跨工况泛化三个方面。在此基础上，报告凝练出四条具体研究空白（失效率基准缺位、失效阈值不一致、多应力耦合机理交互缺乏定量理论、数据驱动模型可解释性弱与泛化差），并逐条对接到学位论文的研究内容与章节，最后展望了多应力耦合定量建模、强物理嵌入混合预测、跨工况泛化与自主化可靠性保证等未来趋势。本报告兼具方法学示范（如何写优秀的研究现状）与领域综述（电容器可靠性现状）双重价值。
"""
    add_body(doc, cn)
    add_heading_cn(doc, "关键词：", level=4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0.82)
    run = p.add_run("电容器可靠性；失效模式；失效机理；可靠性建模；状态监测；剩余使用寿命；"
                    "国内外研究现状；多应力耦合；物理信息神经网络")
    set_cn_font(run, font_name="宋体", size_pt=11)
    add_page_break(doc)

    add_heading_cn(doc, "Abstract", level=0)
    en = """
Capacitors are the most numerous and fundamental passive energy-storage components in power-electronic and modern information systems, and are frequently the weakest link of the host system in terms of field failure rate and service life. A systematic review of the domestic and international research status of capacitor reliability is therefore a core task of electronic-component reliability engineering and of the introduction chapter of a doctoral dissertation. Taking five mainstream families, namely aluminum electrolytic, multilayer ceramic (MLCC), metallized film, solid tantalum and supercapacitors, as the objects, this report deliberately avoids the inefficient styles of chronological listing and of separating domestic and international work into two disjoint parts. Instead, it organizes the review along a single thread of Device - Failure Mode - Failure Mechanism - Reliability Model - Condition Monitoring, providing a contrastive and quantitative review.
The report first establishes the scope, literature-search strategy and organizing thread. It then proceeds thread by thread: at the device level it establishes the failure-attribution status of capacitors as the reliability bottleneck (DC-link capacitors account for about 30 percent of converter failures, electrolytic capacitors 21 to 60 percent); it reviews the open-short-drift taxonomy of failure modes; it analyzes core mechanisms such as electrolyte evaporation (Ea about 0.94 eV), oxygen-vacancy migration (Ea about 0.88 to 1.49 eV), self-healing accumulation (10^4 to 10^6 events at end of life, 5 to 10 percent capacitance loss) and field crystallization (voltage exponent n about 15 to 40); it compares statistical, accelerated/Prokopowicz-Vaskas, generalized Eyring, physics-of-failure and data-driven/PINN hybrid modeling; and it compares on-line C/ESR identification (accuracy improved from about C 4 percent/ESR 12 percent to C 1.5 percent/ESR 5 percent) and machine-learning RUL prediction (RMSE reduced from about 12 percent to about 3.8 percent).
Through comprehensive quantitative comparison, bibliometrics, technology-readiness assessment and a research-evolution timeline, the report concludes that the international and domestic communities exhibit a complementary pattern in which international groups lead paradigm shifts while domestic groups deepen engineering practice, with gaps concentrated in original theory, strongly physics-embedded modeling and cross-condition generalization. Four concrete research gaps are distilled and mapped onto the dissertation research contents, and future trends are outlined. The report serves both as a methodological demonstration of how to write an excellent research-status review and as a substantive review of capacitor reliability.

Keywords: capacitor reliability; failure mode; failure mechanism; reliability modeling; condition monitoring; remaining useful life; research status; multi-stress coupling; physics-informed neural network.
"""
    add_body(doc, en)
    add_page_break(doc)


def build_glossary(doc):
    """缩略语与术语表（不使用编号表格，避免影响全局表号）"""
    add_heading_cn(doc, "缩略语与主要术语", level=0)
    add_body(doc, """
为便于阅读，现将本报告频繁出现的缩略语与主要术语集中说明如下（按英文字母与主题归类）。正文首次出现时一般给出中英文对照，此处提供统一索引。
""")
    items = [
        ("ALT", "Accelerated Life Test，加速寿命试验"),
        ("BME", "Base-Metal-Electrode，贱金属电极（镍内电极、铜端电极MLCC工艺）"),
        ("BOPP", "Biaxially-Oriented PolyPropylene，双向拉伸聚丙烯（薄膜电容介质）"),
        ("C", "Capacitance，电容量（核心健康指标之一）"),
        ("CTE", "Coefficient of Thermal Expansion，热膨胀系数"),
        ("DCL", "DC Leakage current，直流漏电流（钽电容核心退化量）"),
        ("DfR", "Design for Reliability，可靠性内建设计"),
        ("EIS", "Electrochemical/Electrical Impedance Spectroscopy，阻抗谱"),
        ("EoL", "End of Life，寿命终点（失效阈值判据）"),
        ("ESL", "Equivalent Series Inductance，等效串联电感"),
        ("ESR", "Equivalent Series Resistance，等效串联电阻（核心健康指标）"),
        ("FMEA/FMMEA", "失效模式（机理）与影响分析"),
        ("HALT/HAST", "高加速寿命/温湿度应力试验"),
        ("HI", "Health Indicator，健康指标"),
        ("IR", "Insulation Resistance，绝缘电阻（MLCC核心退化量）"),
        ("LSTM/CNN", "长短期记忆网络/卷积神经网络（数据驱动RUL方法）"),
        ("MLCC", "Multilayer Ceramic Capacitor，多层陶瓷电容器"),
        ("PHM", "Prognostics and Health Management，故障预测与健康管理"),
        ("PINN", "Physics-Informed Neural Network，物理信息神经网络"),
        ("PoF", "Physics-of-Failure，失效物理"),
        ("P-V方程", "Prokopowicz-Vaskas方程，陶瓷电容温度—电压寿命经验式"),
        ("RUL", "Remaining Useful Life，剩余使用寿命"),
        ("TC", "Temperature Cycling，温度循环"),
        ("THB", "Temperature-Humidity-Bias，温湿度偏压试验"),
        ("TRL", "Technology Readiness Level，技术成熟度"),
        ("tanδ", "Loss tangent，损耗角正切（介质损耗指标）"),
        ("Ea", "Activation energy，激活能（温度敏感性参数）"),
        ("n", "电压加速指数（逆幂律/P-V方程中的电压指数）"),
        ("β / η", "Weibull分布的形状参数 / 尺度（特征寿命）参数"),
        ("Mission profile", "任务剖面，真实运行工况的应力时间序列"),
        ("Self-healing", "自愈，金属化薄膜电容缺陷点电极蒸发隔离机制"),
        ("Field crystallization", "场致结晶，钽电容Ta2O5非晶向晶态转变"),
    ]
    for k, v in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(2.2)
        p.paragraph_format.first_line_indent = Cm(-2.2)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("%s　" % k)
        set_cn_font(run, font_name="黑体", size_pt=10.5, bold=True)
        run2 = p.add_run(v)
        set_cn_font(run2, font_name="宋体", size_pt=10.5)
    add_page_break(doc)
